"""
API + Export for Shift-based Financial Reports.
"""
from datetime import date, timedelta
from io import BytesIO

from django.http import HttpResponse
from rest_framework import permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView

from core.permissions import get_user_scope
from shifts.report_services import build_all_reports, get_shift_closing_queryset, get_hub_executive_summary


def _parse_date_range(period: str) -> tuple[date, date]:
    """Parse period: daily, weekly, monthly."""
    today = date.today()
    if period == "daily":
        return today, today
    if period == "weekly":
        start = today - timedelta(days=6)
        return start, today
    if period == "monthly":
        start = today.replace(day=1)
        return start, today
    return today, today


class HubExecutiveSummaryView(APIView):
    """ملخص تنفيذي للـ Hub – من إقفالات معتمدة ومعمّدة فقط."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        date_from_str = request.query_params.get("date_from")
        date_to_str = request.query_params.get("date_to")
        date_from = None
        date_to = None
        if date_from_str:
            try:
                date_from = date.fromisoformat(date_from_str)
            except ValueError:
                pass
        if date_to_str:
            try:
                date_to = date.fromisoformat(date_to_str)
            except ValueError:
                pass
        rows = get_hub_executive_summary(
            date_from=date_from,
            date_to=date_to,
            finalized_only=True,
        )
        return Response({"rows": rows})


class ShiftFinancialReportsView(APIView):
    """GET all 5 reports from single source (ShiftClosing)."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        date_from_str = request.query_params.get("date_from")
        date_to_str = request.query_params.get("date_to")
        period = request.query_params.get("period", "weekly")  # daily | weekly | monthly
        branch_ids = request.query_params.get("branch_ids")
        brand_slug = request.query_params.get("brand")
        employee_id = request.query_params.get("employee_id")

        if date_from_str and date_to_str:
            try:
                date_from = date.fromisoformat(date_from_str)
                date_to = date.fromisoformat(date_to_str)
            except ValueError:
                return Response({"detail": "Invalid date format"}, status=400)
        else:
            date_from, date_to = _parse_date_range(period)

        branch_id_list = None
        if branch_ids:
            try:
                branch_id_list = [int(x) for x in branch_ids.split(",") if x.strip()]
            except ValueError:
                pass

        emp_id = None
        if employee_id:
            try:
                emp_id = int(employee_id)
            except ValueError:
                pass

        scope = get_user_scope(request.user)

        reports = build_all_reports(
            date_from=date_from,
            date_to=date_to,
            branch_ids=branch_id_list,
            brand_slug=brand_slug or None,
            employee_id=emp_id,
            scope=scope,
        )

        return Response({
            "date_from": str(date_from),
            "date_to": str(date_to),
            "period": period,
            "reports": reports,
        })


class ShiftFinancialReportExportView(APIView):
    """Export report as Excel or PDF."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        report_type = request.query_params.get("report_type")  # daily_sales, cash_report, etc.
        format_type = request.query_params.get("format", "xlsx")  # xlsx | pdf
        date_from_str = request.query_params.get("date_from")
        date_to_str = request.query_params.get("date_to")
        period = request.query_params.get("period", "weekly")
        branch_ids = request.query_params.get("branch_ids")
        brand_slug = request.query_params.get("brand")

        if not report_type or report_type not in (
            "daily_sales",
            "cash_report",
            "network_report",
            "cashier_shortage",
            "journal_entry",
        ):
            return Response({"detail": "Invalid report_type"}, status=400)

        if date_from_str and date_to_str:
            try:
                date_from = date.fromisoformat(date_from_str)
                date_to = date.fromisoformat(date_to_str)
            except ValueError:
                return Response({"detail": "Invalid date format"}, status=400)
        else:
            date_from, date_to = _parse_date_range(period)

        branch_id_list = None
        if branch_ids:
            try:
                branch_id_list = [int(x) for x in branch_ids.split(",") if x.strip()]
            except ValueError:
                pass

        scope = get_user_scope(request.user)
        reports = build_all_reports(
            date_from=date_from,
            date_to=date_to,
            branch_ids=branch_id_list,
            brand_slug=brand_slug or None,
            scope=scope,
        )
        data = reports.get(report_type, [])

        if format_type == "xlsx":
            buf = _export_excel(report_type, data, date_from, date_to)
            resp = HttpResponse(
                buf.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            resp["Content-Disposition"] = f'attachment; filename="report_{report_type}_{date_from}_{date_to}.xlsx"'
            return resp
        if format_type == "pdf":
            buf = _export_pdf(report_type, data, date_from, date_to)
            resp = HttpResponse(buf.getvalue(), content_type="application/pdf")
            resp["Content-Disposition"] = f'attachment; filename="report_{report_type}_{date_from}_{date_to}.pdf"'
            return resp
        return Response({"detail": "Unsupported format"}, status=400)


def _export_excel(report_type: str, data: list, date_from: date, date_to: date) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.sheet_view.rightToLeft = True
    tajawal = Font(name="Tajawal", size=11)
    tajawal_bold = Font(name="Tajawal", size=11, bold=True)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="Tajawal", size=11, bold=True, color="FFFFFF")

    if report_type == "daily_sales":
        headers = ["الفرع", "التاريخ", "المبيعات", "الضريبة 15%", "الإجمالي"]
        keys = ["branch_name", "date", "sales", "tax", "total_sales"]
    elif report_type == "cash_report":
        headers = ["الفرع", "التاريخ", "النقد الفعلي", "النقد المتوقع", "الفرق"]
        keys = ["branch_name", "date", "actual_cash", "system_cash", "variance"]
    elif report_type == "network_report":
        headers = ["الفرع", "التاريخ", "مدى", "فيزا", "ماستر", "إجمالي الشبكة", "المتوقع", "الفرق"]
        keys = ["branch_name", "date", "mada", "visa", "master_card", "network_total", "system_network", "variance"]
    elif report_type == "cashier_shortage":
        headers = ["التاريخ", "الفرع", "الوردية", "الموظف", "المتوقع", "الفعلي", "العجز/الزيادة"]
        keys = ["date", "branch_name", "shift_type", "employee_name", "expected_amount", "actual_amount", "variance"]
    else:  # journal_entry
        headers = ["التاريخ", "الوصف", "حساب مدين", "حساب دائن", "مبلغ مدين", "مبلغ دائن"]
        keys = ["date", "description", "debit_account", "credit_account", "debit_amount", "credit_amount"]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for row_idx, row in enumerate(data, 2):
        for col_idx, k in enumerate(keys, 1):
            val = row.get(k, "")
            if isinstance(val, (int, float)) and k in ("sales", "tax", "total_sales", "actual_cash", "system_cash",
                                                     "variance", "mada", "visa", "master_card", "network_total",
                                                     "system_network", "expected_amount", "actual_amount",
                                                     "debit_amount", "credit_amount"):
                val = round(float(val), 2)
            ws.cell(row=row_idx, column=col_idx, value=val).font = tajawal

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _export_pdf(report_type: str, data: list, date_from: date, date_to: date) -> BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm)
    styles = getSampleStyleSheet()
    elements = []

    title = Paragraph(
        f"<b>تقرير - {report_type}</b><br/>من {date_from} إلى {date_to}",
        ParagraphStyle(name="Title", fontName="Helvetica-Bold", fontSize=14, alignment=2),
    )
    elements.append(title)
    elements.append(Spacer(1, 0.5*cm))

    if not data:
        elements.append(Paragraph("لا توجد بيانات", styles["Normal"]))
        doc.build(elements)
        buf.seek(0)
        return buf

    if report_type == "daily_sales":
        headers = ["الفرع", "التاريخ", "المبيعات", "الضريبة", "الإجمالي"]
        keys = ["branch_name", "date", "sales", "tax", "total_sales"]
    elif report_type == "cash_report":
        headers = ["الفرع", "التاريخ", "الفعلي", "المتوقع", "الفرق"]
        keys = ["branch_name", "date", "actual_cash", "system_cash", "variance"]
    elif report_type == "network_report":
        headers = ["الفرع", "التاريخ", "مدى", "فيزا", "ماستر", "الإجمالي", "المتوقع", "الفرق"]
        keys = ["branch_name", "date", "mada", "visa", "master_card", "network_total", "system_network", "variance"]
    elif report_type == "cashier_shortage":
        headers = ["التاريخ", "الفرع", "الوردية", "الموظف", "المتوقع", "الفعلي", "العجز/الزيادة"]
        keys = ["date", "branch_name", "shift_type", "employee_name", "expected_amount", "actual_amount", "variance"]
    else:
        headers = ["التاريخ", "الوصف", "مدين", "دائن", "مبلغ مدين", "مبلغ دائن"]
        keys = ["date", "description", "debit_account", "credit_account", "debit_amount", "credit_amount"]

    table_data = [headers]
    for row in data:
        r = []
        for k in keys:
            val = row.get(k, "")
            if isinstance(val, (int, float)):
                val = f"{float(val):,.2f}"
            r.append(str(val))
        table_data.append(r)

    t = Table(table_data)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.black),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
    ]))
    elements.append(t)
    doc.build(elements)
    buf.seek(0)
    return buf
